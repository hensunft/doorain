"""Безопасное извлечение текста: TXT/MD/HTML/PDF(текстовый слой)/DOCX. Без OCR, без сети, без исполнения JS."""

from __future__ import annotations

import re
import time
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from ...core.config import Settings


class UnsupportedDocument(Exception):
    """Понятная причина отказа: показывается пользователю."""


@dataclass
class Parsed:
    pages: list[str]  # текст по страницам (для не-PDF — одна «страница»)
    mime: str
    warnings: list[str] = field(default_factory=list)
    paginated: bool = False


_SKIP = {"script", "style", "head", "noscript", "iframe", "object", "embed", "template", "svg"}
_BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table", "ul", "ol", "blockquote", "pre"}


class _Text(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIP:
            self._skip += 1
        elif tag in _BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in _BLOCK:
            self.out.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.out.append(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            if enc == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    try:
        return data.decode("cp1251")
    except UnicodeDecodeError as e:
        raise UnsupportedDocument("Не удалось определить кодировку текста (ожидается UTF-8 или Windows-1251)") from e


def sniff(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "docx" if ext == ".docx" else "zip"
    head = data[:2048].lower()
    if ext in (".html", ".htm") or b"<html" in head or b"<!doctype html" in head:
        return "html"
    if ext in (".txt", ".md", ".markdown", ".text") or b"\x00" not in data[:4096]:
        return "md" if ext in (".md", ".markdown") else "txt"
    return "binary"


def parse(path: Path, filename: str, s: Settings) -> Parsed:
    data = path.read_bytes()
    kind = sniff(data, filename)
    if kind == "pdf":
        parsed = _pdf(path, s)
    elif kind == "docx":
        parsed = _docx(path, s)
    elif kind == "html":
        p = _Text()
        p.feed(_decode(data))
        parsed = Parsed(["".join(p.out)], "text/html")
    elif kind in ("txt", "md"):
        parsed = Parsed([_decode(data)], "text/markdown" if kind == "md" else "text/plain")
    else:
        raise UnsupportedDocument("Неподдерживаемый тип файла. Допустимы TXT, Markdown, HTML, PDF с текстовым слоем и DOCX")
    total = sum(len(p) for p in parsed.pages)
    if total > s.max_extracted_chars:
        raise UnsupportedDocument("Извлечённый текст превышает допустимый объём")
    if total < 200 or not re.search(r"[A-Za-zА-Яа-яЁё]{3,}", "".join(parsed.pages)):
        hint = " Похоже на скан без текстового слоя: подготовьте текстовую версию (OCR не входит в MVP)." if kind == "pdf" else ""
        raise UnsupportedDocument("В документе не найден текст, пригодный для индексации." + hint)
    return parsed


def _pdf(path: Path, s: Settings) -> Parsed:
    from pypdf import PdfReader  # noqa: PLC0415
    from pypdf.errors import PyPdfError  # noqa: PLC0415

    try:
        r = PdfReader(str(path))
        if r.is_encrypted:
            raise UnsupportedDocument("PDF защищён паролем")
        if len(r.pages) > s.max_pdf_pages:
            raise UnsupportedDocument(f"В PDF более {s.max_pdf_pages} страниц")
        t0, pages = time.monotonic(), []
        for pg in r.pages:
            pages.append(pg.extract_text() or "")
            if time.monotonic() - t0 > 180:
                raise UnsupportedDocument("Превышено время обработки PDF")
    except PyPdfError as e:
        raise UnsupportedDocument("PDF повреждён или не читается") from e
    return Parsed(pages, "application/pdf", paginated=True)


def _docx(path: Path, s: Settings) -> Parsed:
    limit = s.max_docx_unpacked_mb * 1024 * 1024
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if sum(i.file_size for i in infos) > limit or len(infos) > 5000:
                raise UnsupportedDocument("DOCX превышает допустимый размер после распаковки")
            for i in infos:
                if i.compress_size and i.file_size / i.compress_size > 200 and i.file_size > 5_000_000:
                    raise UnsupportedDocument("DOCX похож на архивную бомбу")
            if "word/document.xml" not in z.namelist():
                raise UnsupportedDocument("Файл не является документом DOCX")
        from docx import Document  # noqa: PLC0415

        doc = Document(str(path))
    except zipfile.BadZipFile as e:
        raise UnsupportedDocument("DOCX повреждён") from e
    lines = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return Parsed(["\n".join(lines)], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
